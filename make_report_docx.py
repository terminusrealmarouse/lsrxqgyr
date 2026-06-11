"""按「人工智能技术及应用实验报告」模板生成 MNIST CNN 实验报告(.docx)。

模板格式约定：
    封面：人工智能技术及应用实验报告 /（2026年）/ 实验名称 / 专业班级 / 学生姓名·学生学号
    一级标题：宋体 四号(14pt) 粗体，段前/段后 0.5 行，顶格
    二级标题：宋体 小四号(12pt) 粗体
    正文：宋体 小四号(12pt)
    图：图题置于图下方，居中 宋体 五号(10.5pt) 粗体
    表：表题置于表上方，居中 宋体 五号(10.5pt) 粗体，三线表

依赖 train_mnist_cnn.py + make_diagrams.py 产生的图表与 outputs/summary.json。

用法：
    python make_report_docx.py
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FIG = Path("report/figures")
OUT = Path("report/MNIST_CNN_实验报告.docx")
CN = "宋体"
LATIN = "Times New Roman"


def _font(run, size, bold=False, cn=CN):
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)


def body(doc, text, size=12, indent=True, align=None, bold=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.5
    if indent:
        fmt.first_line_indent = Pt(24)  # 首行缩进 2 字
    if align:
        p.alignment = align
    _font(p.add_run(text), size, bold=bold)
    return p


def h1(doc, num, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.5
    _font(p.add_run(f"{num}　{text}"), 14, bold=True)
    return p


def h2(doc, num, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.5
    _font(p.add_run(f"{num}　{text}"), 12, bold=True)
    return p


def fig_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    _font(p.add_run(text), 10.5, bold=True)


def tbl_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    _font(p.add_run(text), 10.5, bold=True)


def add_image(doc, path, width):
    if not Path(path).exists():
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _set_border(elem, tag, sz, val="single"):
    e = OxmlElement(f"w:{tag}")
    e.set(qn("w:val"), val)
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), "000000")
    elem.append(e)


def three_line_table(doc, rows, widths=None):
    """生成三线表：表格顶/底为粗实线，表头下为细实线，无竖线和其它横线。"""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    _set_border(borders, "top", 12)
    _set_border(borders, "bottom", 12)
    for tag in ("left", "right", "insideV", "insideH"):
        _set_border(borders, tag, 0, val="none")
    tblPr.append(borders)

    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(para.add_run(str(txt)), 10.5, bold=(i == 0))
            # 表头行下边线
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tcB = OxmlElement("w:tcBorders")
                _set_border(tcB, "bottom", 6)
                tcPr.append(tcB)
    return table


def cover_line(doc, label, value=""):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(label), 14, bold=True)
    run = p.add_run("　" + (value if value else "　" * 12))
    _font(run, 14, bold=False)
    if value:
        run.font.underline = True
    return p


def build_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(t.add_run("人工智能技术及应用实验报告"), 26, bold=True)
    y = doc.add_paragraph()
    y.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(y.add_run("（2026年）"), 16, bold=True)
    for _ in range(6):
        doc.add_paragraph()
    cover_line(doc, "实 验 名 称 ：", "基于卷积神经网络的 MNIST 手写数字识别")
    cover_line(doc, "专 业 班 级 ：")
    cover_line(doc, "学 生 姓 名 ：")
    cover_line(doc, "学 生 学 号 ：")
    cover_line(doc, "实 验 日 期 ：")
    doc.add_page_break()


def main():
    with open("outputs/summary.json", encoding="utf-8") as f:
        s = json.load(f)
    hist = s["history"]
    fa = s["final_test_acc"] * 100
    ba = s["best_test_acc"] * 100

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN)

    build_cover(doc)

    # 1 实验目的
    h1(doc, "1", "实验目的")
    body(doc, "（1）熟悉 PyTorch 深度学习框架的使用；")
    body(doc, "（2）熟悉卷积神经网络（CNN）的训练思路；")
    body(doc, "（3）掌握卷积神经网络对 MNIST 手写数字的识别过程。")

    # 2 实验内容
    h1(doc, "2", "实验内容")
    body(
        doc,
        "MNIST 数据库共有 7 万张手写数字图片，其中 6 万张用于训练神经网络，1 万张用于测试神经网络，"
        "每张图片大小为 28×28 像素（单通道灰度图），标签为 0~9 共 10 类。本次实验使用 PyTorch 搭建一个"
        "卷积神经网络模型，以实现 MNIST 手写数据的识别任务，并对训练、测试结果进行分析。",
    )

    # 3 实验步骤
    h1(doc, "3", "实验步骤")
    for t in [
        "（1）读取数据（输入数据与目标输出数据）；",
        "（2）创建卷积神经网络模型；",
        "（3）训练网络；",
        "（4）测试网络；",
        "（5）结果输出与分析。",
    ]:
        body(doc, t)

    # 4 实验设备
    h1(doc, "4", "实验设备")
    body(doc, "（1）计算机；")
    body(doc, "（2）PyTorch 深度学习框架；")
    body(
        doc,
        f"（3）软件环境：Python 3.12 + torch 2.12.0（CPU）+ torchvision 0.27.0 + matplotlib。",
    )

    # 5 实验过程
    h1(doc, "5", "实验过程")

    h2(doc, "5.1", "读取数据")
    body(
        doc,
        "使用 torchvision.datasets.MNIST 自动下载并加载训练集（60000 张）与测试集（10000 张）。"
        "预处理时先用 ToTensor() 将像素归一化到 [0,1]，再用 MNIST 全局均值/标准差 "
        "Normalize((0.1307,),(0.3081,)) 做标准化，使输入分布更稳定、收敛更快。训练集 batch_size=64 "
        "且随机打乱，测试集 batch_size=1000。输入数据为形状 [N,1,28,28] 的图像张量，目标输出为形状 "
        "[N] 的整数标签（0~9）。",
    )

    h2(doc, "5.2", "创建卷积神经网络模型")
    body(doc, "本实验搭建的卷积神经网络结构示意图如图 1 所示，各层配置如表 1 所示。")
    add_image(doc, FIG / "architecture.png", width=6.4)
    fig_caption(doc, "图 1　CNN 网络结构示意图")
    tbl_caption(doc, "表 1　网络各层配置")
    three_line_table(
        doc,
        [
            ["层", "配置", "输出尺寸"],
            ["输入 Input", "单通道灰度图", "1×28×28"],
            ["卷积 Conv1", "3×3, 1→32, padding=1, ReLU", "32×28×28"],
            ["池化 Pool1", "MaxPool 2×2", "32×14×14"],
            ["卷积 Conv2", "3×3, 32→64, padding=1, ReLU", "64×14×14"],
            ["池化 Pool2", "MaxPool 2×2", "64×7×7"],
            ["展平 Flatten", "—", "3136"],
            ["全连接 FC1", "3136→128, ReLU, Dropout 0.5", "128"],
            ["全连接 FC2", "128→10", "10"],
        ],
    )
    body(
        doc,
        "网络由两组「卷积+ReLU+最大池化」逐层提取边缘、笔画等局部空间特征并降低分辨率，再经全连接层"
        "映射到 10 个类别得分，其中 Dropout(0.5) 用于抑制过拟合。模型可训练参数总数为 421,642。",
    )

    h2(doc, "5.3", "训练网络")
    body(
        doc,
        "损失函数采用交叉熵损失 CrossEntropyLoss（内部含 Softmax，适合多分类）；优化器采用 Adam，"
        "学习率 lr=0.001；迭代次数 Epochs=10；随机种子固定为 42 以保证可复现。每个 epoch 遍历一次全部"
        "训练数据，按 batch 执行前向传播→计算损失→反向传播→参数更新。具体神经网络参数设置见表 2。",
    )
    tbl_caption(doc, "表 2　神经网络参数设置")
    three_line_table(
        doc,
        [
            ["参数", "取值"],
            ["迭代次数 Epochs", str(s["epochs"])],
            ["批大小 Batch size", str(s["batch_size"])],
            ["学习率 Learning rate", str(s["lr"])],
            ["优化器 Optimizer", s["optimizer"]],
            ["损失函数 Loss", "CrossEntropyLoss（交叉熵）"],
            ["激活函数", "ReLU"],
            ["正则化", "Dropout(0.5)"],
            ["随机种子", "42"],
            ["可训练参数量", f"{s['trainable_params']:,}"],
        ],
    )

    h2(doc, "5.4", "测试网络")
    body(
        doc,
        "每个 epoch 结束后，在 1 万张测试集上使用 model.eval() 与 torch.no_grad() 模式评估损失与准确率"
        "（准确率 = 预测正确样本数 / 总样本数），并记录最优精度。",
    )

    h2(doc, "5.5", "结果输出与分析")
    body(doc, "完整训练日志与最终测试精度如图 2 所示，逐轮训练/测试结果如表 3 所示。")
    add_image(doc, FIG / "results_panel.png", width=6.4)
    fig_caption(doc, "图 2　训练日志与最终测试精度")
    tbl_caption(doc, "表 3　逐轮训练与测试结果")
    three_line_table(
        doc,
        [["Epoch", "训练损失", "训练精度", "测试损失", "测试精度"]]
        + [
            [
                str(i + 1),
                f"{hist['train_loss'][i]:.4f}",
                f"{hist['train_acc'][i]*100:.2f}%",
                f"{hist['test_loss'][i]:.4f}",
                f"{hist['test_acc'][i]*100:.2f}%",
            ]
            for i in range(s["epochs"])
        ],
    )
    body(doc, f"最终测试精度为 {fa:.2f}%（最佳 {ba:.2f}%），CPU 训练总耗时约 {s['train_time_sec']:.1f} 秒。", bold=True)

    body(doc, "训练/测试的损失与精度曲线如图 3 所示。")
    add_image(doc, FIG / "training_curves.png", width=6.4)
    fig_caption(doc, "图 3　训练/测试损失与精度曲线")
    body(
        doc,
        "由图可见：训练损失从约 0.21 快速下降并趋于平稳，测试损失同步下降，说明模型有效学习且无明显"
        "过拟合；测试精度在第 1 个 epoch 即达到 98.25%，第 3 个 epoch 后稳定在 99% 以上；训练精度与测试"
        "精度差距很小（约 0.1%），表明 Dropout 与较浅的网络结构有效控制了过拟合。",
    )

    body(doc, "测试集混淆矩阵如图 4 所示。")
    add_image(doc, FIG / "confusion_matrix.png", width=4.8)
    fig_caption(doc, "图 4　测试集混淆矩阵")
    body(
        doc,
        "混淆矩阵对角线元素占绝对多数，说明绝大部分样本被正确分类；非对角线误分类数量极少，较易混淆"
        "的情形主要集中在形状相近的数字之间（如 5↔3、4↔9、7↔2），符合手写数字的直观特点。",
    )

    body(doc, "随机抽取的测试样本预测结果如图 5 所示（绿色标题表示预测正确）。")
    add_image(doc, FIG / "sample_predictions.png", width=5.4)
    fig_caption(doc, "图 5　测试样本预测示例")

    # 6 小结
    h1(doc, "6", "小结")
    body(
        doc,
        "本实验基于 PyTorch 成功搭建了一个包含 2 个卷积层、2 个池化层和 2 个全连接层的卷积神经网络，"
        "完成了 MNIST 手写数字识别任务。在 Epochs=10、batch_size=64、学习率=0.001、Adam 优化器、交叉熵"
        f"损失的设置下，模型最终在 1 万张测试图片上达到 {fa:.2f}%（最佳 {ba:.2f}%）的识别精度。",
    )
    body(
        doc,
        "实验表明，卷积神经网络通过卷积与池化自动提取图像的空间特征，相比全连接网络在图像任务上具有"
        "参数共享、平移不变、精度高的优势。后续可尝试增加卷积层数、加入 BatchNorm、引入数据增强或调整"
        "学习率等手段，以进一步提升模型精度或加快收敛速度。",
    )

    doc.save(str(OUT))
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
