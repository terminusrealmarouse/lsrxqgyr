"""按学校模板《人工智能技术及应用实验报告》格式生成 Word 报告。

格式规范(取自模板):
- 封面: 标题 汉仪大宋简 42pt 居中, (2026年) 18pt; 实验名称/专业班级/姓名/学号 宋体四号粗体。
- 一级栏目标题: 宋体 四号(14pt) 粗体, 段前段后 0.5 行。
- 正文: 宋体 小四号(12pt), 1.5 倍行距。
- 二级标题(1.1): 宋体 小四号 粗体。
- 图: 图题置于图下方, 居中, 宋体 五号(10.5pt) 粗体, 编号"图1"。
- 表: 三线表, 表题置于表上方, 居中, 宋体 五号 粗体, 编号"表1"。

用法: python build_report_template_docx.py  ->  人工智能技术及应用实验报告.docx
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, "results")
OUT = os.path.join(HERE, "人工智能技术及应用实验报告.docx")

# 字号(pt)
SZ_TITLE = 42      # 封面大标题
SZ_YEAR = 18
SZ_H1 = 14         # 四号
SZ_BODY = 12       # 小四
SZ_CAP = 10.5      # 五号(图表题)


def set_run(run, text=None, font="宋体", size=SZ_BODY, bold=False):
    if text is not None:
        run.text = text
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    return run


def para(doc, text="", font="宋体", size=SZ_BODY, bold=False, align=None,
         line_spacing=1.5, space_before=None, space_after=None, indent_chars=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line_spacing is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_spacing
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if indent_chars and align in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
        # 正文首行缩进 2 字符
        pf.first_line_indent = Pt(size * indent_chars)
    if text:
        set_run(p.add_run(), text, font=font, size=size, bold=bold)
    return p


def heading1(doc, text):
    """一级栏目: 宋体四号粗体, 段前段后 0.5 行。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(SZ_H1 * 0.5)
    pf.space_after = Pt(SZ_H1 * 0.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    set_run(p.add_run(), text, font="宋体", size=SZ_H1, bold=True)
    return p


def heading2(doc, text):
    """二级标题: 宋体小四粗体。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(SZ_BODY * 0.5)
    pf.space_after = Pt(SZ_BODY * 0.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    set_run(p.add_run(), text, font="宋体", size=SZ_BODY, bold=True)
    return p


def body(doc, text):
    return para(doc, text, font="宋体", size=SZ_BODY, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def figure(doc, filename, caption, width_in=5.3):
    path = os.path.join(RESULTS, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(path):
        p.add_run().add_picture(path, width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    set_run(cap.add_run(), caption, font="宋体", size=SZ_CAP, bold=True)


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))  # 1/8 pt
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


def three_line_table(doc, caption, header, rows):
    """三线表: 表题在上(居中 宋体五号粗体), 仅顶线/表头下线/底线。"""
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    set_run(cap.add_run(), caption, font="宋体", size=SZ_CAP, bold=True)

    n_rows = 1 + len(rows)
    table = doc.add_table(rows=n_rows, cols=len(header))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 表头
    for j, h in enumerate(header):
        c = table.cell(0, j)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(c.paragraphs[0].add_run(), str(h), font="宋体", size=SZ_CAP, bold=True)
    # 数据
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = table.cell(i, j)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(c.paragraphs[0].add_run(), str(val), font="宋体", size=SZ_CAP, bold=False)

    last = n_rows - 1
    for i in range(n_rows):
        for j in range(len(header)):
            top = 12 if i == 0 else None                       # 顶线(粗)
            bottom = None
            if i == 0:
                bottom = 6                                     # 表头下线
            if i == last:
                bottom = 12                                    # 底线(粗)
            _set_cell_border(table.cell(i, j), top=top, bottom=bottom, left=None, right=None)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    for line in code.splitlines():
        set_run(p.add_run(), line + "\n", font="Consolas", size=9, bold=False)


def cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(), "人工智能技术及应用实验报告", font="汉仪大宋简", size=SZ_TITLE, bold=False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(), "（2026年）", font="汉仪大宋简", size=SZ_YEAR, bold=False)

    for _ in range(4):
        doc.add_paragraph()

    labels = [
        "实 验 名 称 ：  基于人工神经网络的相位提取（PyTorch）",
        "专 业 班 级：",
        "学 生 姓 名：                  学 生 学 号：",
    ]
    for text in labels:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 2.0
        p.paragraph_format.left_indent = Pt(48)
        set_run(p.add_run(), text, font="宋体", size=SZ_H1, bold=True)

    doc.add_page_break()


def main():
    doc = Document()
    # 默认正文样式
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(SZ_BODY)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    cover(doc)

    # 一、实验目的
    heading1(doc, "实验目的")
    body(doc, "1. 熟悉 PyTorch 深度学习框架的基本使用；")
    body(doc, "2. 熟悉人工神经网络（多层感知机 MLP）的结构与原理；")
    body(doc, "3. 掌握神经网络的训练过程（前向传播、损失计算、反向传播与参数更新）。")

    # 二、实验内容
    heading1(doc, "实验内容")
    body(doc, "在测量信号中，相位 x 与强度信号 y 之间存在理论对应关系。为方便生成数据集，"
              "假设数据符合如下函数关系：")
    eq = doc.add_paragraph()
    eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(eq.add_run(), "y = 0.5 + 0.3·cos(x) + ε ,   x ∈ (0, π)", font="宋体", size=SZ_BODY, bold=False)
    set_run(eq.add_run(), "                （1）", font="宋体", size=SZ_BODY, bold=False)
    body(doc, "其中 ε 为噪声分布（本实验取高斯噪声 ε ~ N(0, 0.02²)）。要求建立人工神经网络，"
              "求解 x 与 y 的映射模型：输入强度信号 y，即可提取出相位 x 的分布，并观察与分析实验结果。")
    body(doc, "需要说明映射的可行性：余弦函数 cos(x) 在区间 (0, π) 上严格单调递减，因此 y→x 是一一对应"
              "（单射）的，反映射 x = f(y) 存在且唯一，可用神经网络拟合；若区间扩大到 (0, 2π)，"
              "同一个 y 会对应两个 x，映射不再唯一，网络将无法学到确定的反函数。")

    # 三、实验步骤
    heading1(doc, "实验步骤")
    body(doc, "1. 导入（生成）数据：生成输入数据 y 与目标输出数据 x；")
    body(doc, "2. 创建神经网络模型；")
    body(doc, "3. 训练、测试网络；")
    body(doc, "4. 应用网络并输出结果。")

    # 四、实验设备
    heading1(doc, "实验设备")
    body(doc, "1. 计算机；")
    body(doc, "2. PyTorch 框架（CPU 版即可），辅以 NumPy、Matplotlib。")

    # 五、实验过程
    heading1(doc, "实验过程")

    heading2(doc, "1.1 数据生成")
    body(doc, "在 (0, π) 内均匀采样 N = 2000 个相位 x（避开端点 0 与 π），按式(1)计算带噪强度 y；"
              "网络输入为强度 y，目标输出为相位 x；并按 8:2 随机划分训练集（1600）与测试集（400）。"
              "数据集分布及理论曲线如图1所示。")
    figure(doc, "01_dataset.png", "图1  数据集分布与理论曲线 y = 0.5 + 0.3cos(x)")

    heading2(doc, "1.2 网络模型与参数设置")
    body(doc, "采用多层感知机（MLP），结构为 Linear(1,64)→Tanh→Linear(64,64)→Tanh→Linear(64,1)，"
              "输入维度 1（强度 y），输出维度 1（相位 x），隐藏层各 64 个神经元，使用 Tanh 激活，"
              "可训练参数量为 4353。主要参数设置见表1。")
    three_line_table(
        doc,
        "表1  主要参数设置",
        ["参数", "取值", "说明"],
        [
            ["样本数 N_SAMPLES", "2000", "数据集总样本数"],
            ["噪声标准差 NOISE_STD", "0.02", "高斯噪声 ε 的标准差"],
            ["测试集比例 TEST_RATIO", "0.2", "训练:测试 = 8:2"],
            ["网络结构", "1-64-64-1", "MLP，Tanh 激活"],
            ["损失函数", "MSE", "均方误差"],
            ["优化器", "Adam", "—"],
            ["学习率 LR", "1e-3", "—"],
            ["批大小 BATCH_SIZE", "64", "—"],
            ["训练轮数 EPOCHS", "300", "—"],
            ["随机种子 SEED", "42", "保证结果可复现"],
        ],
    )

    heading2(doc, "1.3 PyTorch 核心代码")
    body(doc, "完整代码见 phase_retrieval.py，核心片段如下：")
    code_block(
        doc,
        "# 1. 数据生成: y = 0.5 + 0.3*cos(x) + eps, x in (0, pi)\n"
        "def generate_data(n_samples, noise_std):\n"
        "    x = np.random.uniform(1e-3, np.pi - 1e-3, size=(n_samples, 1))\n"
        "    eps = np.random.normal(0.0, noise_std, size=(n_samples, 1))\n"
        "    y = 0.5 + 0.3 * np.cos(x) + eps\n"
        "    return x.astype(np.float32), y.astype(np.float32)\n\n"
        "# 2. 网络模型: 输入强度 y, 输出相位 x\n"
        "class PhaseNet(nn.Module):\n"
        "    def __init__(self, hidden=64):\n"
        "        super().__init__()\n"
        "        self.net = nn.Sequential(\n"
        "            nn.Linear(1, hidden), nn.Tanh(),\n"
        "            nn.Linear(hidden, hidden), nn.Tanh(),\n"
        "            nn.Linear(hidden, 1))\n"
        "    def forward(self, x):\n"
        "        return self.net(x)\n\n"
        "# 3. 训练: Adam + MSE\n"
        "criterion = nn.MSELoss()\n"
        "optimizer = torch.optim.Adam(model.parameters(), lr=1e-3)\n"
        "for epoch in range(EPOCHS):\n"
        "    for yb, xb in train_loader:        # 注意: 输入 y, 目标 x\n"
        "        optimizer.zero_grad()\n"
        "        loss = criterion(model(yb), xb)\n"
        "        loss.backward(); optimizer.step()\n\n"
        "# 4. 应用: 输入强度 y, 提取相位 x\n"
        "with torch.no_grad():\n"
        "    x_pred = model(y_test_t)\n",
    )

    heading2(doc, "1.4 训练过程")
    body(doc, "训练与测试损失同步快速下降并收敛，二者基本重合，没有出现过拟合，"
              "最终 train MSE ≈ 0.0145，test MSE ≈ 0.0135，损失曲线如图2所示。")
    figure(doc, "02_loss.png", "图2  训练 / 测试损失曲线")

    heading2(doc, "1.5 测试与结果")
    body(doc, "在测试集上用训练好的网络提取相位，平均绝对误差 MAE ≈ 0.0906 rad，"
              "均方根误差 RMSE ≈ 0.1164 rad。预测相位与真实相位散点紧贴 y = x 对角线（图3），"
              "网络学到的反映射 x = f(y) 为一条平滑单调曲线，与真实样本走势一致（图4）。")
    figure(doc, "03_pred_vs_true.png", "图3  测试集 预测相位 vs 真实相位")
    figure(doc, "04_inverse_mapping.png", "图4  网络学到的反映射 x = f(y)")
    body(doc, "给定若干强度值进行相位提取，并与解析解 x = arccos((y-0.5)/0.3) 对比，见表2。")
    three_line_table(
        doc,
        "表2  相位提取应用示例",
        ["输入强度 y", "网络预测 x (rad)", "解析解 x (rad)"],
        [["0.200", "2.9092", "3.1416"], ["0.500", "1.5784", "1.5708"], ["0.800", "0.2111", "0.0000"]],
    )

    heading2(doc, "1.6 结果分析（个人见解）")
    body(doc, "(1) 端点误差较大：当 y 接近极值（x→0 或 x→π）时预测误差明显偏大。原因是 "
              "dy/dx = -0.3sin(x) 在端点处趋于 0，曲线近乎水平，强度对相位不敏感，"
              "同样大小的噪声会被反映射放大成较大的相位误差，这是问题本身的病态性，并非网络缺陷。")
    body(doc, "(2) 中间区域（x ≈ π/2）精度最高，此处 |dy/dx| 最大，映射对噪声最不敏感。")
    body(doc, "(3) 未过拟合：训练/测试损失曲线几乎重合，模型容量与数据规模匹配良好。")
    body(doc, "(4) 可改进方向：增大样本量或降低噪声可整体降低误差；端点附近可通过加密采样"
              "或加权损失缓解误差放大。")

    # 六、小结
    heading1(doc, "小结")
    body(doc, "本实验用 PyTorch 搭建并训练了一个多层感知机，成功实现了从强度 y 到相位 x 的反映射模型。"
              "由于 cos(x) 在 (0, π) 上单调，y→x 映射唯一，网络能够稳定收敛并准确提取相位"
              "（测试集 RMSE ≈ 0.12 rad）。误差主要集中在曲线端点，源于映射在该处的病态性"
              "（灵敏度趋零）与噪声放大。实验完整覆盖了“数据导入→建模→训练测试→应用输出”的"
              "人工神经网络工作流程，达到了熟悉 PyTorch、理解人工神经网络与训练过程的实验目的。")

    doc.save(OUT)
    print("[OK] template report docx saved.")


if __name__ == "__main__":
    main()
